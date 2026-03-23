#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
直接列出demo目录中的所有文件并读取
"""

import os

# 先看看demo目录中到底有什么文件
demo_dir = r"d:\pythonProject\pytest_sxp\outputs\demo"

print("demo目录内容:")
for f in os.listdir(demo_dir):
    fpath = os.path.join(demo_dir, f)
    size = os.path.getsize(fpath)
    print(f"  {f} - {size} bytes")
    
    # 尝试用另一种方式读取
    if f.endswith('.docx'):
        print(f"  尝试读取: {f}")
        
        try:
            from docx import Document
            doc = Document(fpath)
            
            print(f"    段落数: {len(doc.paragraphs)}")
            print(f"    表格数: {len(doc.tables)}")
            
            # 打印段落
            print("    段落内容:")
            for p in doc.paragraphs:
                if p.text.strip():
                    print(f"      {p.text.strip()[:100]}")
            
            # 打印表格
            print("    表格内容:")
            for t_idx, table in enumerate(doc.tables):
                print(f"      Table {t_idx + 1}: {len(table.rows)} rows x {len(table.columns)} cols")
                for r_idx, row in enumerate(table.rows):
                    row_text = [c.text.strip() for c in row.cells]
                    print(f"        Row {r_idx}: {row_text}")
                    
        except Exception as e:
            print(f"    读取错误: {e}")
