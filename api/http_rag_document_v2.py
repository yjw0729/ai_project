#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
RAG文档处理API接口 - 扩展版本
支持PRD和技术文档同时上传和向量化
"""

import os
import json
import asyncio
import uuid
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Any, Optional

from flask import Blueprint, request, jsonify, current_app
from werkzeug.utils import secure_filename

# 导入RAG服务
try:
    from common.rag.services.rag_services import RAGService
    from common.rag.core.models import DocumentType
    from common.rag.core.word_document_processor import WordDocumentProcessor
except ImportError as e:
    current_app.logger.error(f"RAG模块导入失败: {e}")
    RAGService = None

# 创建蓝图
rag_document_opt_v2 = Blueprint('rag_document_opt_v2', __name__)

# 全局RAG服务实例
rag_service = None

# 配置
UPLOAD_FOLDER = 'uploads/rag_docs'
ALLOWED_EXTENSIONS = {'docx', 'doc', 'pdf', 'txt', 'md', 'markdown'}
MAX_FILE_SIZE = 50 * 1024 * 1024  # 50MB

# 确保上传目录存在
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# PRD和技术文档的提示词模板
PROMPT_TEMPLATES = {
    "prd": """你是一个专业的测试工程师和产品经理。请根据以下产品需求文档(PRD)内容，分析并生成全面的测试案例。

## 产品需求文档内容
{content}

## 要求
1. 分析PRD，提取关键功能点、业务流程和用户场景
2. 为每个功能点生成测试案例
3. 测试案例必须包含：
   - title: 测试案例标题
   - module: 功能模块名称
   - priority: 优先级 (P0/P1/P2/P3)
   - test_type: 测试类型 (功能测试/边界测试/异常测试/性能测试)
   - precondition: 前置条件
   - steps: 测试步骤列表
   - expected_result: 预期结果

4. 请以JSON数组格式返回

请生成测试案例：""",
    
    "tech_spec": """你是一个专业的测试工程师和技术文档分析师。请根据以下技术规格文档内容，分析并生成全面的测试案例。

## 技术规格文档内容
{content}

## 要求
1. 分析技术规格文档，提取接口定义、数据结构、业务逻辑
2. 重点关注接口测试、异常处理、数据验证
3. 为每个接口和关键逻辑生成测试案例
4. 测试案例必须包含：
   - title: 测试案例标题
   - module: 功能模块名称
   - priority: 优先级 (P0/P1/P2/P3)
   - test_type: 测试类型 (功能测试/边界测试/异常测试/性能测试)
   - precondition: 前置条件
   - steps: 测试步骤列表
   - expected_result: 预期结果

5. 请以JSON数组格式返回

请生成测试案例：""",
    
    "both": """你是一个专业的测试工程师。请根据以下需求文档（包含PRD和技术规格）内容，全面分析并生成测试案例。

## 需求文档内容
{content}

## 要求
1. 综合分析PRD和技术规格，提取完整的业务功能、接口定义和数据流
2. 生成覆盖业务场景和技术实现的测试案例
3. 测试案例必须包含：
   - title: 测试案例标题
   - module: 功能模块名称
   - priority: 优先级 (P0/P1/P2/P3)
   - test_type: 测试类型 (功能测试/边界测试/异常测试/性能测试)
   - precondition: 前置条件
   - steps: 测试步骤列表
   - expected_result: 预期结果

4. 请以JSON数组格式返回

请生成测试案例："""
}


def get_rag_service_v2():
    """获取RAG服务实例"""
    global rag_service
    if rag_service is None and RAGService is not None:
        try:
            base_dirs = [
                os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                os.getcwd(),
            ]
            
            config_dir = None
            for base_dir in base_dirs:
                test_path = os.path.join(base_dir, "app", "config", "rag", "business_modules.json")
                if os.path.exists(test_path):
                    config_dir = os.path.join(base_dir, "app", "config")
                    break
            
            if config_dir is None:
                config_dir = os.path.abspath("app/config")
            
            rag_service = RAGService(config_dir=config_dir)
            current_app.logger.info("RAG服务初始化成功")
        except Exception as e:
            current_app.logger.error(f"RAG服务初始化失败: {e}")
            return None
    return rag_service


def allowed_file(filename):
    """检查文件类型是否允许"""
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def save_upload_file(file, document_id):
    """保存上传的文件"""
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_ext = filename.rsplit('.', 1)[1].lower()
        new_filename = f"{document_id}.{file_ext}"
        file_path = os.path.join(UPLOAD_FOLDER, new_filename)
        file.save(file_path)
        return file_path, new_filename
    return None, None


def extract_document_content(file_path: str, file_ext: str) -> str:
    """提取文档内容"""
    try:
        if file_ext in ['docx', 'doc']:
            from docx import Document
            doc = Document(file_path)
            content_parts = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    content_parts.append(text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        content_parts.append(" | ".join(row_text))
            return "\n\n".join(content_parts)
        
        elif file_ext in ['pdf']:
            try:
                import PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    text_parts = [page.extract_text() for page in reader.pages]
                    return "\n\n".join([t for t in text_parts if t])
            except:
                return "PDF内容提取失败"
        
        elif file_ext in ['txt', 'md', 'markdown']:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        
        return ""
    except Exception as e:
        current_app.logger.error(f"提取文档内容失败: {e}")
        return ""


def process_document_to_chunks(file_path: str, metadata: dict) -> List:
    """使用WordDocumentProcessor处理文档"""
    processor = WordDocumentProcessor()
    chunks = processor.process_word_document(file_path, metadata=metadata)
    # 过滤有效chunks
    valid_chunks = [c for c in chunks if c.content and 50 <= len(c.content.strip()) <= 2000]
    return valid_chunks


@rag_document_opt_v2.route('/upload_dual', methods=['POST'])
def upload_dual_documents():
    """
    同时上传PRD和技术文档并进行向量化处理
    
    请求参数 (multipart/form-data):
    - prd_file: PRD文档文件 (可选)
    - tech_spec_file: 技术规格文档文件 (可选)
    - collection_name: 集合名称 (可选，默认: "documents")
    - business_module: 业务模块 (必填)
    - document_title: 文档标题 (可选)
    
    注意: prd_file和tech_spec_file至少要上传一个
    
    返回：
    {
        "code": 200,
        "message": "上传成功",
        "data": {
            "prd_document_id": "uuid",
            "tech_spec_document_id": "uuid",
            "collection_name": "documents",
            "business_module": "cross_border_opening",
            "status": "completed",
            "prd_chunks_count": 10,
            "tech_spec_chunks_count": 5,
            "total_chunks_count": 15,
            "prompt_type": "prd|tech_spec|both"
        }
    }
    """
    try:
        rag = get_rag_service_v2()
        if not rag:
            return jsonify({
                "code": 500,
                "message": "RAG服务初始化失败",
                "data": None
            }), 500
        
        # 检查文件
        prd_file = request.files.get('prd_file')
        tech_spec_file = request.files.get('tech_spec_file')
        
        if not prd_file and not tech_spec_file:
            return jsonify({
                "code": 400,
                "message": "至少需要上传一个文档(prd_file或tech_spec_file)",
                "data": None
            }), 400
        
        # 获取参数
        business_module = request.form.get('business_module', '')
        collection_name = request.form.get('collection_name', 'documents')
        document_title = request.form.get('document_title', 'Dual Document')
        
        # 验证业务模块
        if not business_module:
            return jsonify({
                "code": 400,
                "message": "必须指定业务模块 (business_module)",
                "data": None
            }), 400
        
        if not rag.validate_business_module(business_module):
            available_modules = list(rag.get_business_modules().keys())
            return jsonify({
                "code": 400,
                "message": f"无效的业务模块: {business_module}。可用的模块: {', '.join(available_modules)}",
                "data": None
            }), 400
        
        # 处理PRD文档
        prd_document_id = None
        prd_chunks_count = 0
        prd_content = ""
        
        if prd_file and prd_file.filename:
            if not allowed_file(prd_file.filename):
                return jsonify({
                    "code": 400,
                    "message": f"不支持的PRD文件类型: {prd_file.filename}",
                    "data": None
                }), 400
            
            prd_document_id = str(uuid.uuid4())
            prd_file_path, prd_saved_filename = save_upload_file(prd_file, prd_document_id)
            
            if prd_file_path:
                # 提取内容用于后续生成测试案例
                prd_content = extract_document_content(prd_file_path, prd_saved_filename.rsplit('.', 1)[1])
                
                # 向量化
                prd_chunks = process_document_to_chunks(prd_file_path, {
                    "document_id": prd_document_id,
                    "document_type": "product_requirement",
                    "business_module": business_module,
                    "document_title": f"{document_title}_PRD"
                })
                
                if prd_chunks:
                    # 构建索引
                    from common.rag.core.vector_indexer import VectorIndexer
                    indexer = VectorIndexer(rag.config_manager)
                    
                    loop = asyncio.new_event_loop()
                    asyncio.set_event_loop(loop)
                    
                    # 添加到集合
                    collection_name_prd = f"{collection_name}_prd"
                    stats = loop.run_until_complete(
                        indexer.build_index(prd_chunks, collection_name_prd)
                    )
                    prd_chunks_count = len(prd_chunks)
                    loop.close()
                    
                    current_app.logger.info(f"PRD文档向量化完成: {prd_chunks_count} chunks")
        
        # 处理技术规格文档
        tech_spec_document_id = None
        tech_spec_chunks_count = 0
        tech_spec_content = ""
        
        if tech_spec_file and tech_spec_file.filename:
            if not allowed_file(tech_spec_file.filename):
                return jsonify({
                    "code": 400,
                    "message": f"不支持的技术规格文件类型: {tech_spec_file.filename}",
                    "data": None
                }), 400
            
            tech_spec_document_id = str(uuid.uuid4())
            tech_spec_file_path, tech_spec_saved_filename = save_upload_file(tech_spec_file, tech_spec_document_id)
            
            if tech_spec_file_path:
                # 提取内容
                tech_spec_content = extract_document_content(tech_spec_file_path, tech_spec_saved_filename.rsplit('.', 1)[1])
                
                # 向量化
                tech_spec_chunks = process_document_to_chunks(tech_spec_file_path, {
                    "document_id": tech_spec_document_id,
                    "document_type": "technical_specification",
                    "business_module": business_module,
                    "document_title": f"{document_title}_TechSpec"
                })
                
                if tech_spec_chunks:
                    from common.rag.core.vector_indexer import VectorIndexer
                    indexer = VectorIndexer(rag.config_manager)
                    
                    loop = asyncio.new_event_loop()
                    asyncio.set_event_loop(loop)
                    
                    collection_name_tech = f"{collection_name}_tech"
                    stats = loop.run_until_complete(
                        indexer.build_index(tech_spec_chunks, collection_name_tech)
                    )
                    tech_spec_chunks_count = len(tech_spec_chunks)
                    loop.close()
                    
                    current_app.logger.info(f"技术规格文档向量化完成: {tech_spec_chunks_count} chunks")
        
        # 确定提示词类型
        prompt_type = "both"
        if prd_content and not tech_spec_content:
            prompt_type = "prd"
        elif tech_spec_content and not prd_content:
            prompt_type = "tech_spec"
        
        # 合并内容存储（用于后续生成测试案例）
        combined_content = prd_content + "\n\n" + tech_spec_content if (prd_content and tech_spec_content) else (prd_content or tech_spec_content)
        
        # 保存文档内容到临时文件，供后续生成测试案例使用
        if combined_content:
            content_file_path = os.path.join(UPLOAD_FOLDER, f"{document_title}_content.json")
            with open(content_file_path, 'w', encoding='utf-8') as f:
                json.dump({
                    "prd_content": prd_content,
                    "tech_spec_content": tech_spec_content,
                    "combined_content": combined_content,
                    "prompt_type": prompt_type,
                    "business_module": business_module,
                    "document_title": document_title
                }, f, ensure_ascii=False, indent=2)
        
        return jsonify({
            "code": 200,
            "message": "文档上传并处理成功",
            "data": {
                "prd_document_id": prd_document_id,
                "tech_spec_document_id": tech_spec_document_id,
                "collection_name": collection_name,
                "business_module": business_module,
                "status": "completed",
                "prd_chunks_count": prd_chunks_count,
                "tech_spec_chunks_count": tech_spec_chunks_count,
                "total_chunks_count": prd_chunks_count + tech_spec_chunks_count,
                "prompt_type": prompt_type,
                "content_file": os.path.basename(content_file_path) if combined_content else None
            }
        }), 200
    
    except Exception as e:
        current_app.logger.error(f"上传接口异常: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({
            "code": 500,
            "message": f"服务器内部错误: {str(e)}",
            "data": None
        }), 500


@rag_document_opt_v2.route('/generate_test_cases_v2', methods=['POST'])
def generate_test_cases_v2():
    """
    根据已上传的PRD和技术文档生成测试案例
    
    请求参数 (JSON):
    {
        "content_file": "文档内容文件名",  // 上传接口返回的content_file
        "prompt_type": "auto|prd|tech_spec|both",  // 可选，默认auto
        "business_module": "cross_border_opening",  // 可选
        "output_format": "xmind|json"  // 可选，默认xmind
    }
    
    返回：XMind文件下载 或 JSON
    """
    try:
        from common.rag.utils.xmind_generator import XMindGenerator, parse_llm_response_to_test_cases
        
        data = request.get_json()
        if not data:
            return jsonify({
                "code": 400,
                "message": "请求参数不能为空",
                "data": None
            }), 400
        
        content_file = data.get('content_file')
        if not content_file:
            return jsonify({
                "code": 400,
                "message": "缺少content_file参数",
                "data": None
            }), 400
        
        content_file_path = os.path.join(UPLOAD_FOLDER, content_file)
        if not os.path.exists(content_file_path):
            return jsonify({
                "code": 404,
                "message": "文档内容文件不存在，请先上传文档",
                "data": None
            }), 404
        
        # 读取文档内容
        with open(content_file_path, 'r', encoding='utf-8') as f:
            content_data = json.load(f)
        
        combined_content = content_data.get('combined_content', '')
        prompt_type = data.get('prompt_type', 'auto')
        
        if prompt_type == 'auto':
            prompt_type = content_data.get('prompt_type', 'both')
        
        # 选择提示词模板
        prompt_template = PROMPT_TEMPLATES.get(prompt_type, PROMPT_TEMPLATES["both"])
        prompt = prompt_template.format(content=combined_content[:8000])
        
        # 调用LLM生成测试案例
        try:
            rag = get_rag_service_v2()
            
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            
            # 使用RAG服务的LLM客户端
            answer = loop.run_until_complete(
                rag._generate_answer(
                    prompt=prompt,
                    temperature=0.7,
                    max_tokens=4000
                )
            )
            loop.close()
            
            test_cases = parse_llm_response_to_test_cases(answer)
            
        except Exception as e:
            current_app.logger.error(f"LLM调用失败: {e}")
            test_cases = []
        
        if not test_cases:
            # 使用默认模板
            test_cases = [
                {
                    "title": f"验证{content_data.get('document_title', '功能')}基本功能",
                    "module": content_data.get('business_module', '功能验证'),
                    "priority": "P0",
                    "test_type": "功能测试",
                    "precondition": "系统正常运行",
                    "steps": ["打开功能页面", "执行基本操作", "验证结果"],
                    "expected_result": "功能正常"
                }
            ]
        
        output_format = data.get('output_format', 'xmind')
        
        if output_format == 'json':
            return jsonify({
                "code": 200,
                "message": "生成成功",
                "data": {
                    "test_cases": test_cases,
                    "prompt_type": prompt_type
                }
            }), 200
        else:
            # 生成XMind文件
            xmind_gen = XMindGenerator(output_dir="outputs/test_cases")
            xmind_path = xmind_gen.generate_test_cases_xmind(
                document_title=content_data.get('document_title', 'Test Cases'),
                test_cases=test_cases,
                business_module=content_data.get('business_module', '')
            )
            
            return send_file(
                xmind_path,
                as_attachment=True,
                download_name=os.path.basename(xmind_path),
                mimetype='application/octet-stream'
            )
    
    except Exception as e:
        current_app.logger.error(f"生成测试案例失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({
            "code": 500,
            "message": f"生成测试案例失败: {str(e)}",
            "data": None
        }), 500


# 导入send_file
from flask import send_file
