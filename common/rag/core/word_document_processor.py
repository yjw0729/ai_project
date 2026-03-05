#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word文档处理器
专门处理Word格式的产品设计文档，提供智能分割策略
"""

import re
import logging
from typing import List, Dict, Any, Optional, Tuple
from pathlib import Path
from dataclasses import dataclass

from common.rag.core.models import Document, DocumentChunk, DocumentType

logger = logging.getLogger(__name__)

@dataclass
class DocumentSection:
    """文档章节信息"""
    title: str
    content: str
    level: int  # 标题级别 (1-6)
    start_pos: int
    end_pos: int = 0  # 添加默认值
    parent: Optional['DocumentSection'] = None
    children: List['DocumentSection'] = None

    def __post_init__(self):
        if self.children is None:
            self.children = []

class WordDocumentProcessor:
    """Word文档处理器"""

    def __init__(self):
        # 标题识别模式
        self.heading_patterns = [
            # Word文档标题样式
            r'^#{1,6}\s+(.+)$',  # Markdown风格标题
            r'^(\d+\.)\s+(.+)$',  # 数字编号标题 (1. 2. 3.)
            r'^(\d+\.\d+)\s+(.+)$',  # 多级数字编号 (1.1 1.2)
            r'^第[一二三四五六七八九十]+章\s+(.+)$',  # 中文章节标题
            r'^第[一二三四五六七八九十]+节\s+(.+)$',  # 中文小节标题
            r'^(\([0-9]+\))\s+(.+)$',  # 括号编号 (1) (2)
            r'^([A-Z]\.)\s+(.+)$',  # 大写字母编号 (A. B. C.)
            r'^([a-z]\.)\s+(.+)$',  # 小写字母编号 (a. b. c.)
            r'^([IVXLCDM]+\.)\s+(.+)$',  # 罗马数字编号 (I. II. III.)
        ]

        # 产品设计文档专用模式
        self.product_doc_patterns = [
            r'^(产品概述|产品介绍|功能介绍|功能描述)\s*$',
            r'^(技术架构|系统架构|架构设计)\s*$',
            r'^(数据库设计|数据模型)\s*$',
            r'^(API设计|接口设计)\s*$',
            r'^(用户界面|UI设计)\s*$',
            r'^(部署方案|部署架构)\s*$',
            r'^(使用说明|操作指南)\s*$',
            r'^(附录|附件)\s*$',
        ]

    def process_word_document(self, file_path: str, metadata: Dict[str, Any] = None) -> List[DocumentChunk]:
        """
        处理Word文档

        Args:
            file_path: Word文档路径
            metadata: 元数据

        Returns:
            文档块列表
        """
        try:
            # 读取Word文档内容
            content = self._extract_word_content(file_path)

            if not content:
                logger.warning(f"无法提取文档内容: {file_path}")
                return []

            # 智能分割文档
            chunks = self._intelligent_chunk_document(content, metadata or {})

            logger.info(f"Word文档处理完成: {len(chunks)} 个块")
            return chunks

        except Exception as e:
            logger.error(f"处理Word文档失败 {file_path}: {e}")
            return []

    def _extract_word_content(self, file_path: str) -> str:
        """
        提取Word文档内容

        Args:
            file_path: 文档路径

        Returns:
            文档文本内容
        """
        try:
            from docx import Document

            doc = Document(file_path)
            content_parts = []

            # 提取段落内容
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    content_parts.append(text)

            # 提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        content_parts.append(" | ".join(row_text))

            return "\n\n".join(content_parts)

        except ImportError:
            logger.error("未安装python-docx库，无法处理Word文档")
            return ""
        except Exception as e:
            logger.error(f"提取Word内容失败: {e}")
            return ""

    def _intelligent_chunk_document(self, content: str, metadata: Dict[str, Any]) -> List[DocumentChunk]:
        """
        智能分割文档

        Args:
            content: 文档内容
            metadata: 元数据

        Returns:
            文档块列表
        """
        # 1. 按行分割
        lines = content.split('\n')
        cleaned_lines = [line.strip() for line in lines if line.strip()]

        # 2. 识别章节结构
        sections = self._identify_sections(cleaned_lines)

        # 3. 生成文档块
        chunks = []

        for section in sections:
            # 根据章节大小决定分割策略
            if len(section.content) > 2000:  # 长章节进一步分割
                sub_chunks = self._split_long_section(section, metadata)
                chunks.extend(sub_chunks)
            else:
                chunk = self._create_chunk_from_section(section, metadata)
                chunks.append(chunk)

        # 4. 如果没有识别到章节，按段落分割
        if not chunks:
            chunks = self._fallback_chunking(cleaned_lines, metadata)

        return chunks

    def _identify_sections(self, lines: List[str]) -> List[DocumentSection]:
        """
        识别文档章节结构

        Args:
            lines: 文档行列表

        Returns:
            章节列表
        """
        sections = []
        current_section = None
        current_content = []
        current_level = 0

        for i, line in enumerate(lines):
            # 检查是否是标题
            title_info = self._is_heading(line)

            if title_info:
                # 保存之前的章节
                if current_section:
                    current_section.content = '\n'.join(current_content)
                    current_section.end_pos = i - 1
                    sections.append(current_section)

                # 创建新章节
                current_section = DocumentSection(
                    title=title_info['title'],
                    content="",
                    level=title_info['level'],
                    start_pos=i
                )
                current_content = []
                current_level = title_info['level']
            else:
                # 继续积累内容
                if current_section:
                    current_content.append(line)

        # 保存最后一个章节
        if current_section:
            current_section.content = '\n'.join(current_content)
            current_section.end_pos = len(lines) - 1
            sections.append(current_section)

        return sections

    def _is_heading(self, line: str) -> Optional[Dict[str, Any]]:
        """
        检查是否是标题行

        Args:
            line: 文本行

        Returns:
            标题信息或None
        """
        # 清理行内容
        line = line.strip()

        # 检查标题模式
        for pattern in self.heading_patterns:
            match = re.match(pattern, line, re.IGNORECASE)
            if match:
                title = match.groups()[-1]  # 通常标题在最后一个组
                level = self._determine_heading_level(line)
                return {
                    'title': title,
                    'level': level,
                    'pattern': pattern
                }

        # 检查产品文档专用模式
        for pattern in self.product_doc_patterns:
            match = re.match(pattern, line, re.IGNORECASE)
            if match:
                return {
                    'title': line,
                    'level': 2,  # 产品文档章节默认为二级标题
                    'pattern': 'product_doc'
                }

        return None

    def _determine_heading_level(self, line: str) -> int:
        """
        确定标题级别

        Args:
            line: 标题行

        Returns:
            标题级别 (1-6)
        """
        # Markdown风格
        if line.startswith('#'):
            return min(line.count('#'), 6)

        # 中文章节
        if '章' in line:
            return 1
        if '节' in line:
            return 2

        # 数字编号
        if re.match(r'^\d+\.\s', line):
            return 2
        if re.match(r'^\d+\.\d+\s', line):
            return 3

        # 默认级别
        return 2

    def _split_long_section(self, section: DocumentSection, metadata: Dict[str, Any]) -> List[DocumentChunk]:
        """
        分割长章节

        Args:
            section: 章节对象
            metadata: 元数据

        Returns:
            文档块列表
        """
        chunks = []
        content = section.content

        # 按段落分割
        paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]

        current_chunk = ""
        current_tokens = 0
        max_chunk_size = 1500  # 最大块大小

        for paragraph in paragraphs:
            para_tokens = len(paragraph) // 4  # 粗略估算token数量

            if current_tokens + para_tokens > max_chunk_size and current_chunk:
                # 创建块
                chunk = DocumentChunk(
                    content=current_chunk,
                    chunk_id=f"{metadata.get('document_id', 'unknown')}_{section.title}_{len(chunks)}",
                    source_uri=metadata.get('source_uri', ''),
                    metadata={
                        **metadata,
                        'section_title': section.title,
                        'section_level': section.level,
                        'chunk_type': 'section_part',
                        'chunk_index': len(chunks)
                    }
                )
                chunks.append(chunk)

                current_chunk = paragraph
                current_tokens = para_tokens
            else:
                current_chunk += "\n\n" + paragraph if current_chunk else paragraph
                current_tokens += para_tokens

        # 添加最后一个块
        if current_chunk:
            chunk = DocumentChunk(
                content=current_chunk,
                chunk_id=f"{metadata.get('document_id', 'unknown')}_{section.title}_{len(chunks)}",
                source_uri=metadata.get('source_uri', ''),
                metadata={
                    **metadata,
                    'section_title': section.title,
                    'section_level': section.level,
                    'chunk_type': 'section_part',
                    'chunk_index': len(chunks)
                }
            )
            chunks.append(chunk)

        return chunks

    def _create_chunk_from_section(self, section: DocumentSection, metadata: Dict[str, Any]) -> DocumentChunk:
        """
        从章节创建文档块

        Args:
            section: 章节对象
            metadata: 元数据

        Returns:
            文档块
        """
        return DocumentChunk(
            content=f"# {section.title}\n\n{section.content}",
            chunk_id=f"{metadata.get('document_id', 'unknown')}_{section.title}",
            source_uri=metadata.get('source_uri', ''),
            metadata={
                **metadata,
                'section_title': section.title,
                'section_level': section.level,
                'chunk_type': 'section'
            }
        )

    def _fallback_chunking(self, lines: List[str], metadata: Dict[str, Any]) -> List[DocumentChunk]:
        """
        备用分割策略 - 按固定大小分割

        Args:
            lines: 文档行列表
            metadata: 元数据

        Returns:
            文档块列表
        """
        chunks = []
        content = '\n'.join(lines)

        # 按大约1000字符分割
        chunk_size = 1000
        overlap = 200

        start = 0
        chunk_index = 0

        while start < len(content):
            end = start + chunk_size

            # 找到合适的断点（句号、换行等）
            if end < len(content):
                # 向后查找句子结束符
                for i in range(min(100, len(content) - end)):
                    if content[end - i] in '.。!！?？':
                        end = end - i + 1
                        break

            chunk_content = content[start:end].strip()

            if chunk_content:
                chunk = DocumentChunk(
                    content=chunk_content,
                    chunk_id=f"{metadata.get('document_id', 'unknown')}_chunk_{chunk_index}",
                    source_uri=metadata.get('source_uri', ''),
                    metadata={
                        **metadata,
                        'chunk_type': 'fixed_size',
                        'chunk_index': chunk_index,
                        'start_pos': start,
                        'end_pos': end
                    }
                )
                chunks.append(chunk)
                chunk_index += 1

            start = end - overlap

        return chunks

def process_product_design_document(file_path: str, metadata: Dict[str, Any] = None) -> List[DocumentChunk]:
    """
    处理产品设计文档的便捷函数

    Args:
        file_path: 文档路径
        metadata: 元数据

    Returns:
        文档块列表
    """
    processor = WordDocumentProcessor()
    return processor.process_word_document(file_path, metadata or {})




