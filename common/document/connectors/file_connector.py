import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
import docx
import PyPDF2
import markdown
from bs4 import BeautifulSoup
import yaml
import json
from datetime import datetime

from common.document.connectors.base_connector import BaseConnector
from common.document.core.models import Document, DocumentType
from common.document.core.word_document_processor import WordDocumentProcessor


class FileConnector(BaseConnector):
    '''文件系统连接器，支持多种文件格式'''
    def __init__(self, config: Dict[str, Any] = None):
        super().__init__(config)

        self.file_parsers = {
            '.txt': self._parse_text_file,
            '.md': self._parse_markdown_file,
            '.markdown': self._parse_markdown_file,
            '.docx': self._parse_docx_file,
            '.doc': self._parse_docx_file,
            '.pdf': self._parse_pdf_file,
            '.json': self._parse_json_file,
            '.yaml': self._parse_yaml_file,
            '.yml': self._parse_yaml_file,
            '.xml': self._parse_xml_file,
            '.html': self._parse_html_file,
            '.htm': self._parse_html_file,
            '.csv': self._parse_csv_file,
            '.sql': self._parse_sql_file,
            '.py': self._parse_code_file,
            '.java': self._parse_code_file,
            '.js': self._parse_code_file,
            '.ts': self._parse_code_file,
            '.cpp': self._parse_code_file,
            '.c': self._parse_code_file,
            '.go': self._parse_code_file,
            '.rs': self._parse_code_file,
        }

        self.max_file_size = (config or {}).get('max_file_size', 100 * 1024 * 1024) # 最大100M
        self.logger.info(f"文件连接器初始化完成，支持{len(self.file_parsers)}种文件格式")

    def collect(self, source_config: Dict[str, Any]) -> List[Document]:
        """
                从文件系统采集数据

                Args:
                    source_config: 配置参数
                        - paths: 文件路径列表
                        - extensions: 文件扩展名列表（可选）
                        - recursive: 是否递归遍历子目录（默认True）
                        - exclude_patterns: 排除模式列表（正则表达式）
                        - include_patterns: 包含模式列表（正则表达式）

                Returns:
                    文档列表
                """
        paths = source_config.get('paths', [])
        extensions = source_config.get('extensions', list(self.file_parsers.keys()))
        recursive = source_config.get('recursive', True)
        exclude_patterns = source_config.get('exclude_patterns', [])
        include_patterns = source_config.get('include_patterns', [])
        max_depth = source_config.get('max_depth', 10)

        if not paths:
            self.logger.warning("没有指定文件路径")
            return []

        all_documents = []

        for path_str in paths:
            path = Path(path_str).expanduser().resolve()

            if not path.exists():
                self.logger.warning(f"路径不存在:{path}")
                continue

            if path.is_file():
                # 处理单个文件
                if self._should_process_file(path, extensions, exclude_patterns, include_patterns):
                    document = self._process_single_file(path)
                    if document:
                        all_documents.append(document)
            else:
                # 处理目录中的多个文件
                self.logger.info(f"开始扫描目录: {path}")

                file_iterator = path.rglob('*') if recursive else path.glob('*')

                for file_path in file_iterator:
                        if not file_path.is_file():
                            continue

                        # 检查目录深度
                        if recursive and self._get_depth(file_path, path) > max_depth:
                            continue

                        if self._should_process_file(file_path, extensions, exclude_patterns, include_patterns):
                            document = self._process_single_file(file_path)
                            if document:
                                all_documents.append(document)

            self.logger.info(f"从文件系统采集到 {len(all_documents)} 个文档")
            return all_documents

    def _should_process_file(
            self,
            file_path: Path,
            extensions: List[str],
            exclude_patterns: List[str],
            include_patterns: List[str]
    ) -> bool:
        """判断是否应该处理文件"""
        # 检查扩展名 - 支持带点或不带点的扩展名
        if extensions:
            # 标准化扩展名：确保都带点
            normalized_ext = file_path.suffix.lower()
            normalized_extensions = [ext.lower() if ext.startswith('.') else f'.{ext.lower()}' for ext in extensions]
            
            if normalized_ext not in normalized_extensions:
                return False

        # 检查文件大小
        try:
            if file_path.stat().st_size > self.max_file_size:
                self.logger.warning(f"文件过大，跳过: {file_path}")
                return False
        except OSError:
            return False

        # 检查排除模式
        file_path_str = str(file_path)
        for pattern in exclude_patterns:
            if re.search(pattern, file_path_str, re.IGNORECASE):
                return False

        # 检查包含模式
        if include_patterns:
            for pattern in include_patterns:
                if re.search(pattern, file_path_str, re.IGNORECASE):
                    break
            else:
                return False

        return True

    def _get_depth(self, file_path: Path, root_path: Path) -> int:
        """计算文件相对于根目录的深度"""
        return len(file_path.relative_to(root_path).parts)

    def _process_single_file(self, file_path: Path) -> Optional[Document]:
        """处理单个文件"""
        file_extension = file_path.suffix.lower()

        if file_extension not in self.file_parsers:
            self.logger.warning(f"不支持的文件格式: {file_extension} ({file_path})")
            return None

        try:
            # 特殊处理Word文档
            if file_extension in ['.docx', '.doc']:
                # 使用专门的Word文档处理器
                word_processor = WordDocumentProcessor()
                chunks = word_processor.process_word_document(str(file_path))

                if chunks:
                    # 将第一个chunk作为主要内容，其他chunk作为附加信息
                    main_chunk = chunks[0]

                    # 创建Document对象
                    document = Document(
                        content=main_chunk.content,
                        source_uri=str(file_path.absolute()),
                        doc_type=DocumentType.TECH_SPEC,  # 产品设计文档
                        metadata={
                            **main_chunk.metadata,
                            'total_chunks': len(chunks),
                            'all_chunks': [chunk.to_dict() for chunk in chunks]
                        }
                    )
                    return document
                else:
                    self.logger.warning(f"Word文档处理失败: {file_path}")
                    return None

            # 普通文件处理
            parser = self.file_parsers[file_extension]
            content = parser(file_path)

            if not content or not content.strip():
                self.logger.warning(f"文件内容为空: {file_path}")
                return None

            # 获取文件元数据
            stat_info = file_path.stat()
            metadata = {
                'file_path': str(file_path.absolute()),
                'file_name': file_path.name,
                'file_extension': file_extension,
                'file_size': stat_info.st_size,
                'created_time': datetime.fromtimestamp(stat_info.st_ctime).isoformat(),
                'modified_time': datetime.fromtimestamp(stat_info.st_mtime).isoformat(),
                'absolute_path': str(file_path.absolute()),
                'parent_directory': str(file_path.parent),
            }

            # 创建文档对象
            document = self._create_document(
                content=content,
                source_type='file_system',
                source_uri=str(file_path.absolute()),
                metadata=metadata
            )

            self.logger.debug(f"成功处理文件: {file_path}")
            return document

        except Exception as e:
            self.logger.error(f"处理文件失败 {file_path}: {e}")
            return None

    def _parse_text_file(self, file_path: Path) -> str:
        """解析文本文件"""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    def _parse_markdown_file(self, file_path: Path) -> str:
        """解析Markdown文件"""
        content = self._parse_text_file(file_path)

        # 可选：将Markdown转换为纯文本
        try:
            html = markdown.markdown(content)
            soup = BeautifulSoup(html, 'html.parser')
            return soup.get_text(separator='\n')
        except:
            return content

    def _parse_docx_file(self, file_path: Path) -> str:
        """解析Word文档"""
        try:
            doc = docx.Document(file_path)
            paragraphs = [paragraph.text for paragraph in doc.paragraphs]
            tables_content = []

            # 提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    tables_content.append(' | '.join(row_text))

            content = '\n'.join(paragraphs)
            if tables_content:
                content += '\n\n表格内容:\n' + '\n'.join(tables_content)

            return content
        except Exception as e:
            self.logger.warning(f"使用python-docx解析失败，尝试备用方法: {e}")
            return self._parse_text_file(file_path)

    def _parse_pdf_file(self, file_path: Path) -> str:
        """解析PDF文件"""
        try:
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                text_parts = []

                for page_num, page in enumerate(pdf_reader.pages, 1):
                    try:
                        page_text = page.extract_text()
                        if page_text:
                            text_parts.append(f"--- 第 {page_num} 页 ---\n{page_text}")
                    except Exception as e:
                        self.logger.warning(f"提取PDF第{page_num}页失败: {e}")
                        continue

                return '\n\n'.join(text_parts)
        except Exception as e:
            self.logger.error(f"PDF解析失败 {file_path}: {e}")
            return f"[PDF文件，解析失败: {e}]"

    def _parse_json_file(self, file_path: Path) -> str:
        """解析JSON文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return json.dumps(data, ensure_ascii=False, indent=2)
        except Exception as e:
            self.logger.warning(f"JSON解析失败，返回原始内容: {e}")
            return self._parse_text_file(file_path)

    def _parse_yaml_file(self, file_path: Path) -> str:
        """解析YAML文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = yaml.safe_load(f)
                return yaml.dump(data, allow_unicode=True, sort_keys=False)
        except Exception as e:
            self.logger.warning(f"YAML解析失败，返回原始内容: {e}")
            return self._parse_text_file(file_path)

    def _parse_xml_file(self, file_path: Path) -> str:
        """解析XML文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                soup = BeautifulSoup(f.read(), 'xml')
                return soup.get_text()
        except Exception as e:
            self.logger.warning(f"XML解析失败，返回原始内容: {e}")
            return self._parse_text_file(file_path)

    def _parse_html_file(self, file_path: Path) -> str:
        """解析HTML文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                soup = BeautifulSoup(f.read(), 'html.parser')

                # 移除脚本和样式标签
                for script in soup(["script", "style"]):
                    script.decompose()

                # 获取文本
                text = soup.get_text(separator='\n')

                # 清理空白
                lines = (line.strip() for line in text.splitlines())
                chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
                text = '\n'.join(chunk for chunk in chunks if chunk)

                return text
        except Exception as e:
            self.logger.warning(f"HTML解析失败，返回原始内容: {e}")
            return self._parse_text_file(file_path)

    def _parse_csv_file(self, file_path: Path) -> str:
        """解析CSV文件"""
        import csv

        try:
            with open(file_path, 'r', encoding='utf-8', newline='') as f:
                reader = csv.reader(f)
                rows = list(reader)

                if not rows:
                    return ""

                # 如果有标题行
                headers = rows[0] if rows else []
                content_lines = []

                if headers:
                    content_lines.append(" | ".join(headers))
                    content_lines.append("-" * 50)

                for row in rows[1:]:
                    content_lines.append(" | ".join(row))

                return "\n".join(content_lines)
        except Exception as e:
            self.logger.warning(f"CSV解析失败，返回原始内容: {e}")
            return self._parse_text_file(file_path)

    def _parse_sql_file(self, file_path: Path) -> str:
        """解析SQL文件"""
        content = self._parse_text_file(file_path)

        # 美化SQL：添加换行和缩进
        try:
            import sqlparse
            formatted = sqlparse.format(content, reindent=True, keyword_case='upper')
            return formatted
        except ImportError:
            return content

    def _parse_code_file(self, file_path: Path) -> str:
        """解析代码文件"""
        content = self._parse_text_file(file_path)

        # 提取注释和函数/类定义
        lines = content.split('\n')
        important_lines = []

        for line in lines:
            line_stripped = line.strip()

            # 保留注释
            if line_stripped.startswith(('#', '//', '/*', '*', '--')):
                important_lines.append(line)

            # 保留函数/类定义
            elif any(keyword in line for keyword in ['def ', 'class ', 'function ', 'func ', 'pub fn ']):
                important_lines.append(line)

        if important_lines:
            return '\n'.join(important_lines)
        else:
            return content

    def _test_connection_internal(self) -> bool:
        """测试文件系统连接"""
        test_path = Path(".")  # 当前目录
        return test_path.exists() and test_path.is_dir()